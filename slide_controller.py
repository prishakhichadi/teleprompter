from pptx import Presentation

import re

def extract_sentences_from_text(text):
    # very simple sentence splitter for now
    sentences = re.split(r'(?<=[.!?])\s+', text)
    return [s.strip() for s in sentences if len(s.strip()) > 3]


class SlideController:
    def __init__(self, ppt_path: str):
        self.prs = Presentation(ppt_path)
        self.total_slides = len(self.prs.slides)
        self.current_slide = 0  # 0-based index

    def next_slide(self):
        if self.current_slide < self.total_slides - 1:
            self.current_slide += 1

    def prev_slide(self):
        if self.current_slide > 0:
            self.current_slide -= 1

    def get_slide_text(self, idx=None):
        if idx is None:
            idx = self.current_slide

        slide = self.prs.slides[idx]
        text = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)

        return "\n".join(text).strip()
    
    def get_slide_sentences(self, slide_idx):
        text = self.get_slide_text(slide_idx)
        return extract_sentences_from_text(text)



#gives current slide no., slide count, slide text extraction, manual navigation